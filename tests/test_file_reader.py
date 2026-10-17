# SPDX-License-Identifier: GPL-3.0-or-later
"""Unit tests for secure multimodal file reader and document extractor."""

import base64
import io
from pathlib import Path
from unittest.mock import MagicMock

import fitz  # type: ignore
import pytest
from docx import Document  # type: ignore
from openpyxl import Workbook  # type: ignore
from pptx import Presentation  # type: ignore
from pptx.util import Inches, Pt  # type: ignore

from app.db import Database
from app.file_reader import (
    extract_document_pages,
    get_file_bytes,
    ingest_bytes,
    inspect_document_pages,
    parse_page_range,
    read_document,
    search_within_document,
    sweep_blob_cache,
)


@pytest.fixture
def test_db(tmp_path, monkeypatch):
    data_dir = tmp_path / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr("app.config.DATA_DIR", data_dir)
    return Database(data_dir / "doomed.sqlite")


def _make_sample_pdf() -> bytes:
    doc = fitz.open()
    # Page 1: kinematics text
    p1 = doc.new_page()
    p1.insert_text((50, 72), "Kinematics in One Dimension\nVelocity equals displacement over time.\nFormula: v = d / t", fontsize=14)
    # Page 2: with vector drawing
    p2 = doc.new_page()
    p2.insert_text((50, 72), "Velocity Vector Diagram Page", fontsize=14)
    # Add vector drawing shapes
    shape = p2.new_shape()
    shape.draw_line((50, 100), (200, 100))
    shape.draw_line((200, 100), (190, 95))
    shape.draw_line((200, 100), (190, 105))
    shape.draw_rect(fitz.Rect(60, 120, 150, 180))
    shape.finish(color=(1, 0, 0), fill=(0.9, 0.9, 0.9))
    shape.commit()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_sample_docx() -> bytes:
    doc = Document()
    doc.add_heading("Internal Assessment Guide", level=1)
    doc.add_paragraph("Investigation on Simple Harmonic Motion and Pendulum Length.")
    table = doc.add_table(rows=2, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Trial"
    hdr_cells[1].text = "Length (m)"
    hdr_cells[2].text = "Period (s)"
    r1_cells = table.rows[1].cells
    r1_cells[0].text = "1"
    r1_cells[1].text = "0.50"
    r1_cells[2].text = "1.42"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_sample_pptx() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Thermodynamics Lecture"
    slide.placeholders[1].text = "First Law: Energy conservation."
    # Add notes
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = "Remind students that delta U = Q - W."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _make_sample_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Raw Data"
    ws.append(["Mass (kg)", "Force (N)", "Acceleration (m/s^2)"])
    ws.append([2.0, 10.0, 5.0])
    ws.append([4.0, 20.0, 5.0])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_page_range():
    assert parse_page_range("1-3", 10) == [1, 2, 3]
    assert parse_page_range("1, 4, 7-9", 10) == [1, 4, 7, 8, 9]
    assert parse_page_range([2, 5], 10) == [2, 5]
    assert parse_page_range(None, 3) == [1, 2, 3]
    assert parse_page_range(None, 10) == [1, 2, 3, 4, 5]


def test_extract_document_pages():
    pdf_bytes = _make_sample_pdf()
    pages = extract_document_pages(pdf_bytes, "kinematics.pdf")
    assert len(pages) == 2
    assert "Kinematics in One Dimension" in pages[0]["text"]
    assert pages[1]["has_diagrams"] is True  # Detected vector drawing

    docx_bytes = _make_sample_docx()
    d_pages = extract_document_pages(docx_bytes, "ia.docx")
    assert len(d_pages) >= 1
    assert "| Trial | Length (m) | Period (s) |" in d_pages[0]["text"]

    pptx_bytes = _make_sample_pptx()
    p_pages = extract_document_pages(pptx_bytes, "thermo.pptx")
    assert len(p_pages) == 1
    assert "Thermodynamics Lecture" in p_pages[0]["text"]
    assert "delta U = Q - W" in p_pages[0]["text"]

    xlsx_bytes = _make_sample_xlsx()
    x_pages = extract_document_pages(xlsx_bytes, "data.xlsx")
    assert len(x_pages) == 1
    assert "Raw Data" in x_pages[0]["text"]
    assert "| Mass (kg) | Force (N) | Acceleration (m/s^2) |" in x_pages[0]["text"]


def test_private_blob_encryption_at_rest(test_db):
    owner = "tenant1:u1"
    pdf_bytes = _make_sample_pdf()
    res = ingest_bytes(
        db=test_db,
        owner=owner,
        raw_bytes=pdf_bytes,
        filename="secret_notes.pdf",
        source_type="lms",
        ref="lms://resource/10/20",
    )
    sha256 = res["sha256"]
    file_row = test_db.get_file(sha256)
    assert file_row is not None
    assert file_row["encrypted"] == 1

    blob_path = Path(file_row["blob_path"])
    raw_on_disk = blob_path.read_bytes()
    # On disk it MUST be encrypted — raw bytes differ from plaintext
    assert raw_on_disk != pdf_bytes
    assert not raw_on_disk.startswith(b"%PDF")

    # Reading via get_file_bytes for owner decrypts cleanly
    row, decrypted = get_file_bytes(test_db, owner, "lms://resource/10/20")
    assert decrypted == pdf_bytes
    opened_doc = fitz.open(stream=decrypted, filetype="pdf")
    assert "Kinematics in One Dimension" in opened_doc[0].get_text()


def test_per_owner_acl_enforcement(test_db):
    owner_a = "tenant1:student_a"
    owner_b = "tenant1:student_b"
    pdf_bytes = _make_sample_pdf()

    res = ingest_bytes(
        db=test_db,
        owner=owner_a,
        raw_bytes=pdf_bytes,
        filename="student_a_homework.pdf",
        source_type="lms",
        ref="lms://resource/50/60",
    )
    sha = res["sha256"]

    # Student A can read and search
    doc_a = read_document(test_db, owner_a, "lms://resource/50/60")
    assert doc_a["sha256"] == sha
    assert len(doc_a["pages"]) == 2

    # Student B MUST be blocked with PermissionError
    with pytest.raises(PermissionError):
        get_file_bytes(test_db, owner_b, "lms://resource/50/60")

    with pytest.raises(PermissionError):
        get_file_bytes(test_db, owner_b, sha)

    with pytest.raises(PermissionError):
        read_document(test_db, owner_b, sha)


def test_public_ib_blob_access(test_db):
    owner_a = "tenant1:student_a"
    owner_b = "tenant2:student_b"
    pdf_bytes = _make_sample_pdf()

    res = ingest_bytes(
        db=test_db,
        owner=owner_a,
        raw_bytes=pdf_bytes,
        filename="Physics_paper_2_TZ1_HL.pdf",
        source_type="ib",
        ref="ib://fake_sha256",
    )
    sha = res["sha256"]
    file_row = test_db.get_file(sha)
    assert file_row["encrypted"] == 0

    # Public IB file is accessible by any authenticated owner
    row_b, bytes_b = get_file_bytes(test_db, owner_b, sha)
    assert len(bytes_b) == len(pdf_bytes)


def test_search_within_document(test_db):
    owner = "tenant1:u1"
    pdf_bytes = _make_sample_pdf()
    ingest_bytes(
        db=test_db,
        owner=owner,
        raw_bytes=pdf_bytes,
        filename="kinematics.pdf",
        source_type="lms",
        ref="lms://resource/1/1",
    )

    res = search_within_document(test_db, owner, "lms://resource/1/1", "displacement")
    assert len(res["matches"]) >= 1
    assert res["matches"][0]["page_num"] == 1
    assert "displacement" in res["matches"][0]["snippet"].lower()


def test_inspect_document_pages_renders_png(test_db):
    owner = "tenant1:u1"
    pdf_bytes = _make_sample_pdf()
    ingest_bytes(
        db=test_db,
        owner=owner,
        raw_bytes=pdf_bytes,
        filename="kinematics.pdf",
        source_type="lms",
        ref="lms://resource/1/1",
    )

    renders = inspect_document_pages(test_db, owner, "lms://resource/1/1", pages=[2])
    assert len(renders) == 1
    assert renders[0]["page_num"] == 2
    assert renders[0]["mime"] == "image/png"
    assert renders[0]["data_url"].startswith("data:image/png;base64,")

    # Decode and verify PNG header magic bytes
    b64_part = renders[0]["data_url"].split(",", 1)[1]
    png_bytes = base64.b64decode(b64_part)
    assert png_bytes[:8] == b"\x89PNG\r\n\x1a\n"


def test_corrupt_file_handling():
    corrupt_bytes = b"NOT_A_REAL_PDF_DATA_GARBAGE"
    pages = extract_document_pages(corrupt_bytes, "corrupt.pdf")
    assert len(pages) == 1
    assert "Corrupt" in pages[0]["text"]
